"""Üretilen rapor içeriğini (GeneratedReport) düzenlenebilir bir .docx
dosyasına çevirir — png/pdf değil, kullanıcının kendi Word'ünde
düzenleyebileceği gerçek bir belge."""
import io

from docx import Document
from docx.shared import Pt

from ..models.report import GeneratedReport


def build_report_docx(report: GeneratedReport) -> bytes:
    doc = Document()

    title = doc.add_heading(report.title, level=0)
    title.alignment = 1  # ortala

    subtitle = doc.add_paragraph(report.program_name)
    subtitle.alignment = 1
    for run in subtitle.runs:
        run.italic = True
        run.font.size = Pt(11)

    for section in report.sections:
        doc.add_heading(section.heading, level=1)
        for paragraph in section.body.split("\n\n"):
            if paragraph.strip():
                doc.add_paragraph(paragraph.strip())

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
