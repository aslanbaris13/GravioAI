"""Üretilen sunumu (GeneratedPresentation) düzenlenebilir bir .pptx dosyasına
çevirir — PNG değil, kullanıcının kendi PowerPoint'inde düzenleyebileceği
gerçek slaytlar (python-pptx ile üretilen native metin kutuları).

GravioAI'nin kendi marka renkleriyle (lacivert + turuncu) tasarlanmış, boş
varsayılan Office temasının yerine geçen özel bir şablon — düz siyah-beyaz
madde işaretli slaytlar yerine markalı kapak, renkli başlıklar ve vurgulu
madde işaretleri üretir.
"""
import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from ..models.presentation import GeneratedPresentation

NAVY_DARK = RGBColor(0x0A, 0x1F, 0x2D)
ORANGE = RGBColor(0xF9, 0x73, 0x16)
ORANGE_DEEP = RGBColor(0xEA, 0x58, 0x0C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x14, 0x22, 0x2C)
MUTED = RGBColor(0x5A, 0x6B, 0x75)
SUBTLE = RGBColor(0xAE, 0xC4, 0xD2)
FOOTER_LIGHT = RGBColor(0xB9, 0xB3, 0xA6)
FOOTER_DARK = RGBColor(0x5A, 0x70, 0x80)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
_BLANK_LAYOUT = 6


def _set_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _textbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tb.text_frame.word_wrap = True
    tb.text_frame.margin_left = 0
    tb.text_frame.margin_right = 0
    tb.text_frame.margin_top = 0
    tb.text_frame.margin_bottom = 0
    return tb


def _add_text(
    slide, text, x, y, w, h, *, size=14, bold=False, italic=False,
    color=INK, font="Calibri", align=PP_ALIGN.LEFT, anchor=None,
):
    tb = _textbox(slide, x, y, w, h)
    tf = tb.text_frame
    if anchor is not None:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    return tb


def _add_badge(slide, x, y, d, bg: RGBColor, glyph: str, glyph_size=22, glyph_color=WHITE):
    """Marka motifiyle tutarlı, dairesel ikon rozeti (bkz. sunum/uygulama arayüzündeki
    yuvarlak rozet deseni)."""
    shape = slide.shapes.add_shape(9, x, y, d, d)  # MSO_SHAPE.OVAL = 9
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg
    shape.line.fill.background()
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = glyph
    run.font.size = Pt(glyph_size)
    run.font.bold = True
    run.font.color.rgb = glyph_color
    return shape


def _add_bullets(slide, bullets: list[str], x, y, w, h, *, dark=False):
    tb = _textbox(slide, x, y, w, h)
    tf = tb.text_frame
    bullet_color = ORANGE if dark else ORANGE_DEEP
    text_color = WHITE if dark else INK
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(16)
        p.line_spacing = 1.25
        dot = p.add_run()
        dot.text = "●  "
        dot.font.size = Pt(15)
        dot.font.bold = True
        dot.font.color.rgb = bullet_color
        dot.font.name = "Calibri"
        text = p.add_run()
        text.text = bullet
        text.font.size = Pt(15)
        text.font.color.rgb = text_color
        text.font.name = "Calibri"
    return tb


def _page_number(slide, n: int, *, dark=False):
    _add_text(
        slide, f"{n:02d}", Inches(12.55), Inches(7.08), Inches(0.6), Inches(0.3),
        size=10, color=FOOTER_DARK if dark else FOOTER_LIGHT, align=PP_ALIGN.RIGHT,
    )


def build_presentation_pptx(presentation: GeneratedPresentation) -> bytes:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[_BLANK_LAYOUT]

    # ---- Kapak ----
    cover = prs.slides.add_slide(blank)
    _set_bg(cover, NAVY_DARK)
    _add_badge(cover, Inches(0.9), Inches(0.85), Inches(0.8), ORANGE, "◎", 26)
    _add_text(
        cover, presentation.title, Inches(0.9), Inches(2.5), Inches(10.5), Inches(1.8),
        size=40, bold=True, color=WHITE, font="Cambria",
    )
    if presentation.subtitle:
        _add_text(
            cover, presentation.subtitle, Inches(0.9), Inches(4.15), Inches(9.5), Inches(1),
            size=16, italic=True, color=SUBTLE, font="Calibri",
        )
    _add_text(
        cover, "GravioAI ile oluşturuldu", Inches(0.9), Inches(6.75), Inches(6), Inches(0.4),
        size=11, color=FOOTER_DARK, font="Calibri",
    )

    # ---- İçerik slaytları ----
    total = len(presentation.slides)
    for i, slide_data in enumerate(presentation.slides, start=1):
        is_closing = i == total
        slide = prs.slides.add_slide(blank)
        _set_bg(slide, NAVY_DARK if is_closing else WHITE)

        _add_text(
            slide, f"{i:02d}", Inches(0.7), Inches(0.55), Inches(1.2), Inches(0.4),
            size=13, bold=True, color=ORANGE if is_closing else ORANGE_DEEP, font="Calibri",
        )
        _add_text(
            slide, slide_data.heading, Inches(0.7), Inches(0.92), Inches(11.9), Inches(0.9),
            size=30, bold=True, color=WHITE if is_closing else INK, font="Cambria",
        )

        if slide_data.bullets:
            _add_bullets(
                slide, slide_data.bullets, Inches(0.7), Inches(2.1), Inches(11.6), Inches(4.9),
                dark=is_closing,
            )

        _page_number(slide, i + 1, dark=is_closing)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
